"""Konversi AHS_Jakarta_Bekasi.xlsx -> data/ahs.js (lapisan AHS di peta, terpisah dari data outlet).

Koordinat yang tertukar (latitude berisi 106.x, longitude berisi -6.x) dibetulkan otomatis;
baris tanpa koordinat valid dilewati dan dilaporkan. Kota "BEKASI" ditentukan Kota/Kab. dari
kecamatannya (scripts/kecamatan_kota.json berisi kecamatan Kota Bekasi).

Pemakaian: python3 scripts/convert_ahs.py AHS_Jakarta_Bekasi.xlsx
"""
import json
import pathlib
import re
import sys

import openpyxl

from stamp import stamp

ROOT = pathlib.Path(__file__).resolve().parent.parent
KEC = json.loads((ROOT / "scripts" / "kecamatan_kota.json").read_text(encoding="utf-8"))
COLS = ["id", "Nama AHS", "Alamat", "No. HP", "Kecamatan", "Kota", "latitude", "longitude"]


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def tc(v):
    return " ".join(w[:1].upper() + w[1:].lower() for w in s(v).split())


def ok(lat, lng):
    return isinstance(lat, (int, float)) and isinstance(lng, (int, float)) and -7.2 < lat < -5.8 and 106.3 < lng < 107.7


src = sys.argv[1] if len(sys.argv) > 1 else "AHS_Jakarta_Bekasi.xlsx"
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
data = list(ws.iter_rows(values_only=True))
head = [s(h) for h in data[0]]
col = {h: i for i, h in enumerate(head)}
rows, fixed, skipped = [], [], []
for r in data[1:]:
    g = lambda h: r[col[h]] if h in col and col[h] < len(r) else None
    name = s(g("Nama AHS"))
    if not name:
        continue
    lat, lng = g("Latitude"), g("Longitude")
    if not ok(lat, lng) and ok(lng, lat):
        lat, lng = lng, lat
        fixed.append(name)
    if not ok(lat, lng):
        skipped.append(name)
        continue
    kec = tc(g("Kecamatan"))
    kota = s(g("Kota/Kab")).upper()
    hit = KEC.get(re.sub(r"[^a-z]", "", kec.lower()))
    if "JAKARTA" in kota:
        kota = "Kota " + tc(kota)
    elif "BEKASI" in kota:
        kota = "Kota Bekasi" if (kota.startswith("KOTA") or (hit and hit[1] == "Kota Bekasi")) else "Kab. Bekasi"
    else:
        kota = tc(kota)
    no = s(g("No"))
    rows.append(["ahs-" + (no or str(len(rows) + 1)), name, s(g("Alamat")), s(g("No. HP")), kec, kota,
                 round(float(lat), 7), round(float(lng), 7)])

with open(ROOT / "data" / "ahs.js", "w", encoding="utf-8") as f:
    f.write("window.AHS_DATA=")
    json.dump({"cols": COLS, "rows": rows}, f, ensure_ascii=False, separators=(",", ":"))
    f.write(";\n")
print(f"{len(rows)} AHS ditulis ke data/ahs.js")
print(f"koordinat tertukar dibetulkan ({len(fixed)}): {', '.join(fixed)}")
print(f"dilewati, koordinat tidak valid ({len(skipped)}): {', '.join(skipped)}")
stamp()
